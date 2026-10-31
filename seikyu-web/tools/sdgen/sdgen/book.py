# 補足入力の読み書き。1画面＝1ファイル（補足入力\<画面ID>.xlsx）。シートの並び（構成）と中身（表・図・項目・手順）
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from sdgen import cfg

SHEETS = {
    # 画面シートを上から並べた部品。種類：表紙・戻る・見出し大・見出し中・文・空行・表・図・項目・手順
    '構成': ['順', '種類', '区画ID', '文字', '見出し1', '見出し2'],
    '表': ['表', '順', '左', '右'],                       # 表＝構成の「順」
    '図': ['図', 'ラベル', '画像', '幅pt', '矢印の終点', '起点x', '起点y', '終点x', '終点y'],
    '項目': ['区画ID', 'エリア', 'キー', '項目名称', '区分', '必須', '初期値', '桁数', '少数', '形式', '参照元テーブル', '参照元項目', '備考'],
    '手順': ['区画ID', '手順', '利用者の操作', '内容', '画面の反応', '備考'],
}
GUIDE = [
    '補足入力：画面から自動で取れない内容。1ファイル＝1画面シート。',
    '構成シート：画面シートを上から並べた部品。種類＝表紙／戻る／見出し大／見出し中／文／空行（文字＝行数）／表（見出し1・2＝列見出し、中身は表シート）／図（中身は図シート）／項目（区画IDの画面項目。文字＝見出し、{n}＝項目数）／手順（区画IDの操作手順）。',
    '項目シート：区分・必須・初期値・桁数・形式が空欄なら画面から自動で決める（tools/sdgen/sdgen/rules.py）。値を書くとその値を使う。キー＝画面の項目名（見出し・ラベル）、ラベルのない欄は name 属性。',
    '図シート：図＝構成の「順」。画像＝ fixed:ファイル名（tools/sdgen/fixed）／ capture:撮影名（＋で縦に並べる）／ report:帳票名。矢印の終点＝矢印を引く先の図（画面シートの上からの番号）、起点x・y／終点x・y は図の中の位置（0〜1）。',
]


def path_of(sc):
    return cfg.SUPPLEMENT / f'{sc}.xlsx'


def screens():
    return sorted(p.stem for p in cfg.SUPPLEMENT.glob('SC-*.xlsx'))


def load(sc):
    p = path_of(sc); out = {k: [] for k in SHEETS}
    if not p.exists(): return out
    wb = openpyxl.load_workbook(p, data_only=True)
    for name in SHEETS:
        if name not in wb.sheetnames: continue
        ws = wb[name]; head = [c.value for c in ws[1]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not any(v not in (None, '') for v in r): continue
            out[name].append({h: ('' if v is None else v) for h, v in zip(head, r) if h})
    return out


def save(sc, data):
    cfg.SUPPLEMENT.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    g = wb.create_sheet('説明')
    for i, t in enumerate(GUIDE, 1): g.cell(i, 1, t)
    g.column_dimensions['A'].width = 160
    for name, cols in SHEETS.items():
        ws = wb.create_sheet(name); ws.append(cols)
        for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='CADCFC')
        for r in data.get(name, []): ws.append([r.get(c, '') for c in cols])
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 70 if c in ('文字', '右', '備考', '画面の反応', '内容', '利用者の操作') else 16
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes = 'A2'
    wb.save(path_of(sc))
