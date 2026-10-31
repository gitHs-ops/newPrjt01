# 3 設計書の作成：補足入力\<画面ID>.xlsx の構成の順に画面シートを組み立てる。画面項目は画面から取り出した項目（work\capture）＋規則で作る
#   python build.py [出力.xlsx] [画面ID ...]     … 画面IDを省くと補足入力の全画面
import datetime as dt, json, sys
from pathlib import Path
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter as L
from PIL import Image
from sdgen import cfg, book, rules, style as S

sys.stdout.reconfigure(encoding='utf-8')
COLS = ['Seq', '項目名称', '表示 / 区分', '必須', '初期値', '桁数', '少数', '形式', '参照元 / テーブル', '参照元項目', '備考']
FIELDS = ['項目名称', '区分', '必須', '初期値', '桁数', '少数', '形式', '参照元テーブル', '参照元項目', '備考']
LEGEND = '凡例　I＝入力／I/O＝条件付きの入力および表示／O＝表示のみ。水色＝入力項目、薄橙＝必須項目。〔 〕は画面上のエリア（検索条件・入力フォーム・一覧）。'
KIND_NOTE = [('I', '入力項目。この画面で入力する'), ('I/O', '条件付きの入力および表示'), ('O', '表示のみ（一覧・計算結果・採番値）')]
PX = 96 / 72


def cw(c1, c2='K'):
    return sum(S.WIDTH[chr(c)] for c in range(ord(c1), ord(c2) + 1))


def merge_items(sc, sup, report):
    """補足入力の項目（並び・説明）と画面の項目（自動の値）を合わせる → {区画ID: [行]}"""
    f = cfg.WORK / 'capture' / f'{sc}.json'
    cap = json.loads(f.read_text(encoding='utf-8')) if f.exists() else {'views': []}
    if not f.exists(): report.append(f'{sc}: 画面の取り出し結果なし（補足入力の値だけで作成）')
    dom = {}
    for v in cap['views']:
        for a in v['areas']: dom.setdefault((v['id'], a['area']), []).extend(a['items'])
    out = {}
    for r in sup['項目']:
        items = dom.get((r['区画ID'], r['エリア']), [])
        it = next((x for x in items if r['キー'] != '' and str(x.get('key')) == str(r['キー'])), None)
        row = dict(r)
        if it:
            items.remove(it); a = rules.auto(it, r)
            for c in ['区分', '必須', '初期値', '桁数', '形式']:
                if row[c] in ('', None): row[c] = a[c]
        elif cap['views']:
            report.append(f'{sc}: 補足のみ（画面に無い）: {r["区画ID"]} {r["エリア"]} {r["項目名称"]}（キー {r["キー"]}）')
        for c in ['区分', '必須', '初期値', '桁数', '形式']:   # 自動の列だけ「-」で埋める（説明の列の空欄はそのまま）
            if row.get(c) in ('', None): row[c] = '-'
        out.setdefault(r['区画ID'], []).append(row)
    for (vid, area), items in dom.items():   # 補足入力に無い画面の項目は末尾に載せて知らせる
        for it in items:
            a = rules.auto(it); row = {'区画ID': vid, 'エリア': area, '項目名称': it.get('label') or it.get('key'), **a, '少数': '-', '参照元テーブル': '-', '参照元項目': '-', '備考': '（補足入力なし）'}
            out.setdefault(vid, []).append(row); report.append(f'{sc}: 画面のみ（補足入力に無い）: {vid} {area} {row["項目名称"]}')
    return out


class Sheet:
    def __init__(self, ws):
        self.ws, self.r = ws, 1
        for c, w in S.WIDTH.items(): ws.column_dimensions[c].width = w
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = 'landscape'; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0

    def cell(self, col, v, f=None, fl=None, al=S.TOP, box=True, r=None):
        c = self.ws[f'{col}{r or self.r}']; c.value = v
        c.font = f or S.font(); c.alignment = al
        if fl: c.fill = S.fill(fl)
        if box: c.border = S.BOX
        return c

    def merge(self, c1, c2, r=None):
        r = r or self.r; self.ws.merge_cells(f'{c1}{r}:{c2}{r}')
        for x in range(ord(c1), ord(c2) + 1): self.ws[f'{chr(x)}{r}'].border = S.BOX

    def skip(self, n=1): self.r += n

    def heading(self, text, size=10.5, h=None):
        self.cell('A', text, S.font(size, True, S.NAVY), box=False, al=None)
        if h: self.ws.row_dimensions[self.r].height = h
        self.skip()

    def head(self, cols):   # [(列, 見出し, 結合の終わり)]
        for c1, t, c2 in cols:
            self.cell(c1, t, S.font(9, True), 'head', S.CENTER)
            if c2: self.merge(c1, c2)
        self.skip()

    def row(self, cols, h=None):   # [(列, 値, 結合の終わり, 塗り, 中央)]
        hw = []
        for c1, v, c2, fl, center in cols:
            self.cell(c1, v, fl=fl, al=S.CENTER if center else S.TOP)
            if c2: self.merge(c1, c2)
            hw.append((v, cw(c1, c2 or c1)))
        self.ws.row_dimensions[self.r].height = h or S.height(hw); self.skip()

    def image(self, png, width_pt):
        img = Image.open(png); w = width_pt * PX; h = img.height * w / img.width
        x = XLImage(str(png)); x.width, x.height = round(w), round(h); x.anchor = f'A{self.r}'
        self.ws.add_image(x); self.skip(int(h / PX / 15) + 1)


def picture(spec, name):
    """図の指定（fixed:ファイル／capture:撮影名＋撮影名／report:帳票名）を PNG のパスにする"""
    spec = str(spec or '')
    if spec.startswith('fixed:'): return cfg.FIXED / spec[6:]
    if spec.startswith('report:'): return cfg.WORK / 'capture' / 'png' / (Path(spec[7:]).stem + '.png')
    if spec.startswith('capture:'):
        parts = [cfg.WORK / 'capture' / 'png' / f'{p}.png' for p in spec[8:].split('+')]
        if len(parts) == 1: return parts[0]
        if not all(p.exists() for p in parts): return None
        ims = [Image.open(p).convert('RGB') for p in parts]; w = max(i.width for i in ims)
        out = Image.new('RGB', (w, sum(i.height for i in ims) + 24 * (len(ims) - 1)), 'white'); y = 0
        for i in ims: out.paste(i, (0, y)); y += i.height + 24
        f = cfg.WORK / 'capture' / 'png' / f'{name}.png'; out.save(f); return f
    return None


def items_block(s, title, rows, seq):
    ws = s.ws
    s.heading(title.replace('{n}', str(len(rows))), 13, 24); s.skip()
    s.heading('■ 表示区分の内訳'); s.head([('A', '表示区分', None), ('B', '件数', None), ('C', '意味', 'G')])
    count_rows = []
    for k, note in KIND_NOTE: count_rows.append(s.r); s.row([('A', k, None, None, True), ('B', '', None, None, True), ('C', note, 'G', None, False)])
    total = s.r; s.row([('A', '合計', None, None, True), ('B', f'=SUM(B{count_rows[0]}:B{count_rows[-1]})', None, None, True), ('C', 'この画面に表示される項目', 'G', None, False)])
    for c in 'ABC': ws[f'{c}{total}'].font = S.font(9, True)
    s.skip(); s.heading('■ 画面項目')
    s.head([(L(i + 1), t, None) for i, t in enumerate(COLS)]); ws.row_dimensions[s.r - 1].height = 28
    first, area = s.r, None
    for x in rows:
        if x['エリア'] != area:
            area = x['エリア']; s.cell('A', area, S.font(9, True), 'area'); s.merge('A', 'K'); s.skip()
        seq += 1; inp = x['区分'] in ('I', 'I/O')
        s.row([('A', seq, None, None, True), ('B', x['項目名称'], None, 'input' if inp else None, False), ('C', x['区分'], None, None, True), ('D', x['必須'], None, 'req' if x['必須'] == '○' else None, True),
               ('E', x['初期値'], None, None, False), ('F', x['桁数'], None, None, True), ('G', x['少数'], None, None, True), ('H', x['形式'], None, None, False),
               ('I', x['参照元テーブル'], None, None, False), ('J', x['参照元項目'], None, None, False), ('K', x['備考'], None, None, False)])
    last = max(s.r - 1, first)
    for r_, (k, _) in zip(count_rows, KIND_NOTE):
        c = ws[f'B{r_}']; c.value = f'=COUNTIF($C${first}:$C${last},"{k}")'; c.font = S.font(); c.alignment = S.CENTER; c.border = S.BOX
    s.cell('A', LEGEND, S.font(9, color=S.GRAY), box=False); ws.merge_cells(f'A{s.r}:K{s.r}'); s.skip()
    return seq


def build_screen(wb, sc, report, date):
    sup = book.load(sc); items = merge_items(sc, sup, report)
    tables, figs = {}, {int(f['図']): f for f in sup['図']}
    for t in sup['表']: tables.setdefault(int(t['表']), []).append(t)
    ws = wb.create_sheet(sc); s = Sheet(ws); seq = 0; pics, used = [], set()
    for b in sorted(sup['構成'], key=lambda b: int(b['順'])):
        k, txt = b['種類'], b['文字']
        if k == '表紙':
            s.head([('A', '顧客名', None), ('B', 'システム名', 'C'), ('D', '文書', None), ('E', '作成日', None), ('F', '作成者', None)])
            for c in 'ABCDEF': ws[f'{c}{s.r - 1}'].font = S.font(9, True, S.GRAY); ws[f'{c}{s.r - 1}'].fill = S.fill('meta'); ws[f'{c}{s.r - 1}'].alignment = S.TOP
            s.row([('A', S.CUSTOMER, None, None, False), ('B', S.SYSTEM, 'C', None, False), ('D', S.DOCNAME, None, None, False), ('E', date, None, None, False), ('F', '', None, None, False)])
            ws[f'E{s.r - 1}'].number_format = 'yyyy/mm/dd'
        elif k == '戻る':
            c = s.cell('A', '← 目次へ戻る', S.font(9, color=S.BLUE, underline='single'), box=False); c.hyperlink = "#'目次'!A1"; s.skip()
        elif k == '見出し大': s.heading(txt, 13, 24)
        elif k == '見出し中': s.heading(txt)
        elif k == '文': s.cell('A', txt, box=False); s.skip()
        elif k == '空行': s.skip(int(txt or 1))
        elif k == '表':
            num = b['見出し1'] == 'No'
            s.head([('A', b['見出し1'], None), ('B', b['見出し2'], 'K')])
            for t in sorted(tables.get(int(b['順']), []), key=lambda t: int(t['順'])):
                if t['左'] == '' and t['右'] == '': s.skip(); continue
                s.row([('A', t['左'], None, None, num), ('B', t['右'], 'K', None, False)])
        elif k == '図':
            f = figs.get(int(b['順']), {})
            if f.get('ラベル'): s.cell('A', f['ラベル'], S.font(11, True), box=False, al=None); s.skip()
            p = picture(f.get('画像'), f'{sc}_図{b["順"]}')
            if p and p.exists(): s.image(p, float(f.get('幅pt') or 750)); pics.append(f)
            elif f.get('画像'): report.append(f'{sc}: 図が無い: {f.get("ラベル")} {f.get("画像")}')
        elif k == '項目':
            used.add(b['区画ID']); seq = items_block(s, txt, items.get(b['区画ID'], []), seq)
        elif k == '手順':
            s.heading(txt, 13, 24); s.skip()
            s.head([('A', '手順', None), ('B', '利用者の操作', None), ('C', '内容', 'E'), ('F', '画面の反応', 'I'), ('J', '備考', 'K')])
            for x in [x for x in sup['手順'] if x['区画ID'] == b['区画ID']]:
                s.row([('A', x['手順'], None, None, True), ('B', x['利用者の操作'], None, None, False), ('C', x['内容'], 'E', None, False), ('F', x['画面の反応'], 'I', None, False), ('J', x['備考'], 'K', None, False)])
        else: report.append(f'{sc}: 不明な部品 {k}')
    for vid in set(items) - used: report.append(f'{sc}: 構成に「項目」の無い区画 {vid}（{len(items[vid])} 項目）')
    arrows = [{'from': n, 'to': int(f['矢印の終点']), 'x0': f['起点x'], 'y0': f['起点y'], 'x1': f['終点x'], 'y1': f['終点y']}
              for n, f in enumerate(pics, 1) if str(f.get('矢印の終点') or '').strip()]
    return {'sheet': sc, 'arrows': arrows, 'items': sum(len(v) for v in items.values())}


def main(out=None, screens=None):
    scs = screens or book.screens()
    date = dt.datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    wb = openpyxl.Workbook(); toc = wb.active; toc.title = '目次'
    report, meta = [], []
    for sc in scs: meta.append(build_screen(wb, sc, report, date))
    toc['A1'] = 'APLink 請求・入金管理システム　メニュー別機能対応表'; toc['A1'].font = S.font(13, True, S.NAVY)
    for i, sc in enumerate(scs, 3):
        c = toc.cell(i, 1, sc); c.hyperlink = f"#'{sc}'!A1"; c.font = S.font(9, color=S.BLUE, underline='single')
    out = Path(out or cfg.WORK / 'build' / f'画面設計_生成_{date:%Y%m%d}.xlsx'); out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    info = {'out': str(out), 'sheets': meta}
    out.with_suffix('.json').write_text(json.dumps(info, ensure_ascii=False, indent=1), encoding='utf-8')
    (cfg.WORK / 'build').mkdir(parents=True, exist_ok=True)
    (cfg.WORK / 'build' / 'build.json').write_text(json.dumps(info, ensure_ascii=False, indent=1), encoding='utf-8')
    out.with_name(out.stem + '_report.txt').write_text('\n'.join(report) or '指摘なし', encoding='utf-8')
    print(f'{out}  {len(meta)} シート・項目 {sum(m["items"] for m in meta)}'); print('\n'.join(report) or '指摘なし')


if __name__ == '__main__':
    a = sys.argv[1:]
    main(a[0] if a and a[0].endswith('.xlsx') else None, [x for x in a if not x.endswith('.xlsx')] or None)
